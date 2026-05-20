"""Export a kherveDOC Document model tree to a .docx (Word) file.

Reverse of the .docx importer: walks the model and builds a python-docx
Document, mapping each node type back to the Word style / formatting it
originally came from.
"""
from __future__ import annotations

from pathlib import Path
from typing import Callable

from docx import Document as DocxDocument
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from . import model as M

# Type for the optional progress callback: (current_block, total_blocks).
ProgressCallback = Callable[[int, int], None] | None

# ---- alignment mapping ---------------------------------------------------

_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

# ---- heading level → Word style name -------------------------------------

_HEADING_STYLES = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3",
                   4: "Heading 4", 5: "Heading 5"}


# ---- inline rendering ----------------------------------------------------

def _add_inlines(para, inlines: list[M.Inline]) -> None:
    """Append inline nodes to a python-docx Paragraph as runs."""
    for inline in inlines:
        if isinstance(inline, M.Text):
            run = para.add_run(inline.text)
            for mark in inline.marks:
                if mark == "bold":
                    run.bold = True
                elif mark == "italic":
                    run.italic = True
                elif mark == "underline":
                    run.underline = True
                elif mark == "strikethrough":
                    run.font.strike = True
                elif mark == "subscript":
                    run.font.subscript = True
                elif mark == "superscript":
                    run.font.superscript = True
                elif mark == "smallcaps":
                    run.font.small_caps = True
                elif mark == "code":
                    run.font.name = "Courier New"

        elif isinstance(inline, M.MathInline):
            run = para.add_run(inline.latex)
            run.italic = True

        elif isinstance(inline, M.Link):
            _add_hyperlink(para, inline.url, inline.children)

        elif isinstance(inline, M.Footnote):
            text = "".join(
                i.text for i in inline.children if isinstance(i, M.Text))
            run = para.add_run(f"[{text}]")
            run.font.superscript = True
            run.font.size = Pt(8)

        elif isinstance(inline, M.Citation):
            run = para.add_run(f"[{', '.join(inline.keys)}]")

        elif isinstance(inline, M.CrossRef):
            run = para.add_run(f"[{inline.kind}:{inline.label}]")

        elif isinstance(inline, M.InlineRaw):
            run = para.add_run(inline.latex)
            run.font.name = "Courier New"

        elif isinstance(inline, M.Highlight):
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            run = para.add_run()
            _add_inlines_to_run_group(para, inline.children)
            # Apply highlight via XML — python-docx has limited highlight API
            color_map = {
                "yellow": "yellow", "green": "green", "blue": "cyan",
                "pink": "magenta", "orange": "darkYellow",
            }
            wd_color = color_map.get(inline.color, "yellow")
            for r in para.runs[-len(inline.children):]:
                rpr = r._element.get_or_add_rPr()
                hi = OxmlElement("w:highlight")
                hi.set(qn("w:val"), wd_color)
                rpr.append(hi)

        elif isinstance(inline, M.Comment):
            _add_inlines(para, inline.children)
            if inline.note:
                run = para.add_run(f" [{inline.note}]")
                run.italic = True
                run.font.size = Pt(8)


def _add_inlines_to_run_group(para, inlines: list[M.Inline]) -> None:
    """Add inlines as runs (used for nested groups like Highlight children)."""
    _add_inlines(para, inlines)


def _add_hyperlink(para, url: str, children: list[M.Inline]) -> None:
    """Insert a hyperlink into a paragraph using low-level OOXML."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    part = para.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    for inline in children:
        if isinstance(inline, M.Text):
            new_run = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            r_style = OxmlElement("w:rStyle")
            r_style.set(qn("w:val"), "Hyperlink")
            rpr.append(r_style)
            new_run.append(rpr)
            t = OxmlElement("w:t")
            t.text = inline.text
            new_run.append(t)
            hyperlink.append(new_run)
        else:
            new_run = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = inline.text if hasattr(inline, "text") else str(inline)
            new_run.append(t)
            hyperlink.append(new_run)

    para._element.append(hyperlink)


# ---- block rendering -----------------------------------------------------

def export_docx(doc: M.Document, path: Path,
                 progress: ProgressCallback = None) -> None:
    """Serialize a kherveDOC Document to a .docx file at *path*.

    *progress*, if given, is called as ``progress(i, total)`` after each
    block is written so the UI can update a progress bar.
    """
    out = DocxDocument()

    meta = doc.meta

    # Set core properties.
    out.core_properties.title = meta.title
    out.core_properties.author = meta.author

    # Set page margins from meta.
    for section in out.sections:
        section.top_margin = Cm(meta.margin_top_cm)
        section.bottom_margin = Cm(meta.margin_bottom_cm)
        section.left_margin = Cm(meta.margin_left_cm)
        section.right_margin = Cm(meta.margin_right_cm)

    # Remove the default empty paragraph that python-docx inserts.
    if out.paragraphs:
        p = out.paragraphs[0]._element
        p.getparent().remove(p)

    total = len(doc.children)

    for i, block in enumerate(doc.children):
        if isinstance(block, M.Title):
            para = out.add_paragraph(style="Title")
            _add_inlines(para, block.children)

        elif isinstance(block, M.Author):
            para = out.add_paragraph(style="Subtitle")
            _add_inlines(para, block.children)

        elif isinstance(block, M.Abstract):
            # Group consecutive abstracts visually — just add as italic paras.
            para = out.add_paragraph()
            _add_inlines(para, block.children)
            for run in para.runs:
                run.italic = True

        elif isinstance(block, M.Keywords):
            para = out.add_paragraph()
            run = para.add_run("Keywords: ")
            run.bold = True
            _add_inlines(para, block.children)

        elif isinstance(block, M.Section):
            style = _HEADING_STYLES.get(block.level, "Heading 5")
            para = out.add_paragraph(style=style)
            _add_inlines(para, block.children)

        elif isinstance(block, M.Paragraph):
            para = out.add_paragraph()
            _add_inlines(para, block.children)
            para.alignment = _ALIGN_MAP.get(block.alignment,
                                            WD_ALIGN_PARAGRAPH.JUSTIFY)

        elif isinstance(block, M.List):
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                para = out.add_paragraph(style=style)
                _add_inlines(para, item.children)

        elif isinstance(block, M.MathBlock):
            para = out.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(block.latex)
            run.italic = True
            run.font.name = "Cambria Math"

        elif isinstance(block, M.Figure):
            para = out.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fig_path = Path(block.path)
            if fig_path.exists():
                run = para.add_run()
                run.add_picture(str(fig_path), width=Inches(4.5))
            else:
                para.add_run(f"[Image: {block.path}]")
            if block.caption:
                cap = out.add_paragraph(block.caption)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cap.runs:
                    run.italic = True

        elif isinstance(block, M.Table):
            if block.rows:
                tbl = out.add_table(rows=len(block.rows),
                                    cols=len(block.rows[0]),
                                    style="Table Grid")
                for r_idx, row in enumerate(block.rows):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < len(tbl.rows[r_idx].cells):
                            tbl.rows[r_idx].cells[c_idx].text = cell_text
                # Bold the first row (header).
                if len(block.rows) > 1:
                    for cell in tbl.rows[0].cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                if block.caption:
                    cap = out.add_paragraph(block.caption)
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in cap.runs:
                        run.italic = True

        elif isinstance(block, M.RawLatex):
            para = out.add_paragraph()
            run = para.add_run(block.text)
            run.font.name = "Courier New"

        elif isinstance(block, M.Frame):
            para = out.add_paragraph(style="Heading 1")
            _add_inlines(para, block.children)

        if progress is not None:
            progress(i + 1, total)

    out.save(str(path))
