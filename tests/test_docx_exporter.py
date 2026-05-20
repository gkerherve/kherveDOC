"""Tests for the .docx exporter — model → .docx round-trip verification."""

import tempfile
from pathlib import Path

import docx

from khervedoc.docx_exporter import export_docx
from khervedoc.model import (
    Abstract, Author, Document, DocMeta, Figure, Keywords, Link,
    List as ListNode, ListItem, MathBlock, MathInline, Paragraph,
    Section, Table, Text, Title,
)


def _export_and_read(doc: Document) -> docx.Document:
    """Export to a temp .docx and re-open it with python-docx."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = Path(f.name)
    export_docx(doc, path)
    return docx.Document(str(path))


def test_title_and_author():
    doc = Document(children=[
        Title(children=[Text(text="My Paper")]),
        Author(children=[Text(text="Jane Doe")]),
    ])
    out = _export_and_read(doc)
    styles = [p.style.name for p in out.paragraphs]
    assert "Title" in styles
    assert "Subtitle" in styles
    assert out.paragraphs[0].text == "My Paper"
    assert out.paragraphs[1].text == "Jane Doe"


def test_section_levels():
    doc = Document(children=[
        Section(level=1, children=[Text(text="Intro")]),
        Section(level=2, children=[Text(text="Sub")]),
        Section(level=3, children=[Text(text="SubSub")]),
    ])
    out = _export_and_read(doc)
    assert out.paragraphs[0].style.name == "Heading 1"
    assert out.paragraphs[1].style.name == "Heading 2"
    assert out.paragraphs[2].style.name == "Heading 3"


def test_paragraph_with_marks():
    doc = Document(children=[
        Paragraph(children=[
            Text(text="bold", marks=["bold"]),
            Text(text=" normal"),
            Text(text=" italic", marks=["italic"]),
        ]),
    ])
    out = _export_and_read(doc)
    runs = out.paragraphs[0].runs
    assert runs[0].bold is True
    assert runs[0].text == "bold"
    assert runs[2].italic is True


def test_paragraph_alignment():
    doc = Document(children=[
        Paragraph(children=[Text(text="centered")], alignment="center"),
    ])
    out = _export_and_read(doc)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    assert out.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_bullet_list():
    doc = Document(children=[
        ListNode(ordered=False, items=[
            ListItem(children=[Text(text="one")]),
            ListItem(children=[Text(text="two")]),
        ]),
    ])
    out = _export_and_read(doc)
    assert len(out.paragraphs) == 2
    assert out.paragraphs[0].text == "one"
    assert out.paragraphs[1].text == "two"


def test_table_export():
    doc = Document(children=[
        Table(rows=[["Name", "Age"], ["Alice", "30"]], caption="People"),
    ])
    out = _export_and_read(doc)
    assert len(out.tables) == 1
    tbl = out.tables[0]
    assert tbl.rows[0].cells[0].text == "Name"
    assert tbl.rows[1].cells[1].text == "30"


def test_math_block():
    doc = Document(children=[
        MathBlock(latex="E = mc^2"),
    ])
    out = _export_and_read(doc)
    assert "E = mc^2" in out.paragraphs[0].text


def test_inline_math():
    doc = Document(children=[
        Paragraph(children=[
            Text(text="Energy is "),
            MathInline(latex="E=mc^2"),
        ]),
    ])
    out = _export_and_read(doc)
    assert "E=mc^2" in out.paragraphs[0].text


def test_page_margins():
    meta = DocMeta(margin_top_cm=3.0, margin_bottom_cm=3.0,
                   margin_left_cm=2.0, margin_right_cm=2.0)
    doc = Document(meta=meta, children=[
        Paragraph(children=[Text(text="hello")]),
    ])
    out = _export_and_read(doc)
    from docx.shared import Cm
    section = out.sections[0]
    assert abs(section.top_margin - Cm(3.0)) < Cm(0.01)
    assert abs(section.left_margin - Cm(2.0)) < Cm(0.01)


def test_abstract_and_keywords():
    doc = Document(children=[
        Abstract(children=[Text(text="This is the abstract.")]),
        Keywords(children=[Text(text="python, latex")]),
    ])
    out = _export_and_read(doc)
    assert "This is the abstract." in out.paragraphs[0].text
    assert "Keywords:" in out.paragraphs[1].text
    assert "python, latex" in out.paragraphs[1].text


def test_core_properties():
    meta = DocMeta(title="Test Title", author="John Smith")
    doc = Document(meta=meta, children=[
        Paragraph(children=[Text(text="body")]),
    ])
    out = _export_and_read(doc)
    assert out.core_properties.title == "Test Title"
    assert out.core_properties.author == "John Smith"
